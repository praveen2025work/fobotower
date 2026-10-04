"""A case as an Excel workbook — for any capability.

Sheets:
  Items      one row per item: its group, the group's finding and decision, and
             the item columns (manifest `export.columns`, default `items.display`)
  Groups     one row per group: finding, verdict, decision, who, when, ticket
  Case       the case's key, status, version, due date, and who exported it

Built from the same read as the case page (`case_detail`), so it shows
exactly what the caller may see — protected fields stay as the page shows them.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from helix.cases import case_detail
from helix.entitlement import Caller
from helix.runner import pinned

HEADER = PatternFill("solid", fgColor="00395D")


def _cell(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        return v
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return str(v)


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEADER
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append([_cell(v) for v in r])
    ws.freeze_panes = "A2"
    for i, h in enumerate(header, start=1):
        width = max([len(str(h))] + [len(str(_cell(r[i - 1]) or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    return ws


async def workbook(case_id: str, caller: Caller) -> tuple[bytes, str]:
    d = await case_detail(case_id, caller)
    from helix.db import get_session
    from helix.models import Case

    async with get_session() as s:
        case = await s.get(Case, case_id)
    m = await pinned(case)
    columns = m.export.columns or d["columns"]
    by_item = {}
    for g in d["groups"]:
        for i in g["item_ids"]:
            by_item[i] = g
    wb = Workbook()

    item_rows = []
    for it in d["items"]:
        g = by_item.get(it["item_id"])
        f, dec = (g or {}).get("finding") or {}, (g or {}).get("decision") or {}
        item_rows.append([it["item_id"], "yes" if it.get("in_scope") else "no", (g or {}).get("label", ""),
                          f.get("status", ""), f.get("verdict", ""), dec.get("action", "pending" if g else ""),
                          dec.get("decided_by", ""), *[it.get(c) for c in columns]])
    _sheet(wb, d["labels"]["item"][:28] + "s", [
        d["labels"]["item"], "in scope", "group", "finding", "verdict", "decision", "decided by", *columns],
        item_rows, first=True)

    group_rows = []
    for g in d["groups"]:
        f, dec, t = g.get("finding") or {}, g.get("decision") or {}, g.get("ticket") or {}
        group_rows.append([g["label"], len(g["item_ids"]), f.get("status", ""), f.get("decided_by", ""),
                           f.get("verdict", ""), f.get("requires_confirmation", ""),
                           f.get("comment") or f.get("reason") or "", dec.get("action", "pending"),
                           dec.get("decided_by", ""), dec.get("decided_at"), dec.get("comment") or "",
                           t.get("reference", "")])
    _sheet(wb, "Groups", ["group", "items", "finding", "proposed by", "verdict", "needs confirmation",
                          "explanation", "decision", "decided by", "decided at", "reviewer comment",
                          "ticket"], group_rows)

    case_rows = [["capability", f"{m.name} ({m.id}) v{d['manifest_version']}"],
                 ["team group", d.get("team_group") or ""], ["case", case_id], ["subject", d["subject"]],
                 *[[k, v] for k, v in d["case_key"].items()],
                 ["status", d["status"]], ["outcome", d.get("outcome") or ""],
                 ["opened", f"{_cell(d['opened_at'])} by {d['opened_by']}"],
                 ["due", _cell(d.get("due_at")) or ""],
                 ["exported by", caller.user_id]]
    _sheet(wb, "Case", ["field", "value"], case_rows)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{case_id}.xlsx"
