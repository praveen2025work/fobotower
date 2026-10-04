"""Review controls and the features around them, each switched on by
configuration (the FOBO groups use them all): business-word group labels,
"Approve all" leaving flagged groups, explicit confirmation, escalation
tickets, deadlines, recurring items, Excel export, delegation and measured
review time."""

from datetime import datetime, timedelta, timezone
from io import BytesIO

import pytest
from openpyxl import load_workbook

from helix import deadlines, gateway
from helix.entitlement import Caller
from tests.helix.conftest import FOBO, RECON

KEY = {"book": "PRIME-MB-04", "cob": "2026-09-30"}


async def _run(api, key=KEY, user="frank"):
    res = await api.post(f"/api/capabilities/{RECON}/cases", headers=api.as_user(user),
                         json={"case_key": key, "team_group": FOBO})
    assert res.status_code == 201, res.text
    return res.json()


def _group(case, gid):
    return next(g for g in case["groups"] if g["group_id"] == gid)


async def _decide(api, case_id, gid, user="frank", **body):
    return await api.post(f"/api/cases/{case_id}/decisions", headers=api.as_user(user),
                          json={"group_id": gid, "action": "approve",
                                "idempotency_key": f"k-{gid}-{user}-{len(body)}", **body})


async def test_groups_are_named_in_business_words(api):
    case = await _run(api)
    labels = {g["group_id"]: g["label"] for g in case["groups"]}
    assert labels == {"C-BO": "Redemption break · back office", "E-FO": "Trade booking break · front office"}


async def test_approve_all_leaves_a_verdict_that_needs_confirmation(api):
    case = await _run(api)
    assert _group(case, "C-BO")["flags"] == ["confirmation"]
    assert _group(case, "C-BO")["bulk_blockers"] == ["confirmation"]
    res = await api.post(f"/api/cases/{case['case_id']}/decisions/bulk", headers=api.as_user("frank"),
                         json={"group_ids": ["C-BO", "E-FO"], "action": "approve",
                               "idempotency_key": "bulk-0001", "review_seconds": 40})
    body = res.json()
    assert [d["group_id"] for d in body["decided"]] == ["E-FO"]
    assert body["refused"][0]["group_id"] == "C-BO"
    assert "confirmation" in body["refused"][0]["reason"]


async def test_confirming_a_flagged_verdict_needs_a_tick_and_words(api):
    case = await _run(api)
    cid = case["case_id"]
    no_tick = await _decide(api, cid, "C-BO", comment="Statement arrived late; post.")
    assert no_tick.status_code == 409 and "confirmation" in no_tick.json()["detail"]
    no_words = await _decide(api, cid, "C-BO", confirmed=True)
    assert no_words.status_code == 409 and "explanation" in no_words.json()["detail"]
    ok = await _decide(api, cid, "C-BO", confirmed=True, comment="Materiality confirmed with PC; post.")
    assert ok.status_code == 201, ok.text
    assert _group(ok.json()["case"], "C-BO")["decision"]["confirmed"] is True


async def test_a_fix_upstream_becomes_one_ticket_for_the_owning_team(api):
    case = await _run(api)
    cid = case["case_id"]
    assert (await _decide(api, cid, "E-FO", review_seconds=30)).status_code == 201
    done = await _decide(api, cid, "C-BO", confirmed=True, comment="Confirmed with PC.", review_seconds=90)
    detail = done.json()["case"]
    assert detail["status"] == "completed"
    ticket = _group(detail, "E-FO")["ticket"]
    assert ticket["status"] == "raised" and ticket["reference"].startswith("INC")
    assert _group(detail, "C-BO")["ticket"] is None             # a POST is no one's fix
    audit = [t for t in detail["tool_calls"] if t["tool"] == "ticketing.create_ticket"]
    assert len(audit) == 1 and audit[0]["requested_by"] == "escalate"
    assert audit[0]["arguments"]["team"] == "Desk"


async def test_only_the_record_step_may_raise_the_capabilitys_ticket():
    who = Caller(user_id="frank", roles=frozenset({"FOBO_CONTROLLER"}), data_scopes={"book": frozenset({"*"})})
    tool = "ticketing.create_ticket"
    args = {"team": "Desk", "title": "t"}
    model = gateway.CallContext(capability_id=RECON, caller=who, allowed_tools=frozenset({tool}),
                                requested_by="llm", write_approved_by="frank", escalation_tool=tool)
    with pytest.raises(gateway.ToolDenied):
        await gateway.call(model, tool, args)
    undecided = gateway.CallContext(capability_id=RECON, caller=who, allowed_tools=frozenset({tool}),
                                    requested_by="escalate", escalation_tool=tool)
    with pytest.raises(gateway.ToolDenied):
        await gateway.call(undecided, tool, args)


async def test_a_case_is_due_the_business_day_after_its_cob_and_reminders_go_once(api):
    case = await _run(api)
    due = datetime.fromisoformat(case["due_at"])
    assert due == datetime(2026, 10, 1, 11, 0, tzinfo=timezone.utc)     # Wed COB → Thu 11:00
    assert await deadlines.check(datetime(2026, 10, 1, 8, 0, tzinfo=timezone.utc)) == []
    assert await deadlines.check(datetime(2026, 10, 1, 9, 30, tzinfo=timezone.utc)) == [(case["case_id"], "soon")]
    assert await deadlines.check(datetime(2026, 10, 1, 9, 45, tzinfo=timezone.utc)) == []
    assert await deadlines.check(datetime(2026, 10, 1, 11, 5, tzinfo=timezone.utc)) == [(case["case_id"], "missed")]
    assert await deadlines.check(datetime(2026, 10, 2, 11, 5, tzinfo=timezone.utc)) == []
    inbox = (await api.get("/api/inbox", headers=api.as_user("frank"))).json()
    row = next(r for r in inbox if r["case_id"] == case["case_id"])
    assert row["due_state"] in ("on_time", "due_soon", "overdue") and row["exposure"] > 0


async def test_an_instrument_breaking_again_on_the_same_book_is_flagged_recurring(api):
    first = await _run(api, {"book": "PRIME-MB-04", "cob": "2026-09-24"})   # EURUSD FWD breaks on both
    second = await _run(api)
    before = {i for g in first["groups"] for i in g["item_ids"]}
    now = {i for g in second["groups"] for i in g["item_ids"]}
    assert before & now                                      # the stub data repeats some breaks
    assert set(second["recurring"]) == before & now
    for item_id in before & now:
        assert second["recurring"][item_id]["runs"] == 2
    other_book = await _run(api, {"book": "PRIME-MB-03", "cob": "2026-09-30"})
    assert other_book["recurring"] == {}                     # recurring is per book
    board = (await api.get(f"/api/capabilities/{RECON}/recurring", headers=api.as_user("frank"))).json()
    assert {r["item_id"] for r in board if r["same"] == {"book": "PRIME-MB-04"}} == before & now


async def test_a_case_downloads_as_a_workbook(api):
    case = await _run(api)
    res = await api.get(f"/api/cases/{case['case_id']}/export.xlsx", headers=api.as_user("frank"))
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Breaks", "Groups", "Case"]
    header = [c.value for c in wb["Breaks"][1]]
    assert header[:3] == ["Break", "in scope", "group"] and "difference" in header
    assert {r[0].value for r in wb["Groups"].iter_rows(min_row=2)} == {
        "Redemption break · back office", "Trade booking break · front office"}
    outsider = await api.get(f"/api/cases/{case['case_id']}/export.xlsx", headers=api.as_user("dan"))
    assert outsider.status_code == 404


async def test_a_colleague_covers_for_an_absent_reviewer_within_their_scope(api):
    case = await _run(api)
    cid = case["case_id"]
    assert (await _decide(api, cid, "E-FO", user="dan")).status_code == 403
    until = (datetime.now(timezone.utc) + timedelta(days=3)).isoformat()
    res = await api.post("/api/me/delegations", headers=api.as_user("frank"),
                         json={"to_user": "dan", "until": until, "reason": "Annual leave"})
    assert res.status_code == 201, res.text
    inbox = (await api.get("/api/inbox", headers=api.as_user("dan"))).json()
    assert any(r["case_id"] == cid and r["acting_for"] == "frank" for r in inbox)
    ok = await _decide(api, cid, "E-FO", user="dan")
    assert ok.status_code == 201, ok.text
    decision = _group(ok.json()["case"], "E-FO")["decision"]
    assert (decision["decided_by"], decision["on_behalf_of"]) == ("dan", "frank")
    mine = (await api.get("/api/me/delegations", headers=api.as_user("frank"))).json()
    await api.delete(f"/api/me/delegations/{mine['away'][0]['delegation_id']}", headers=api.as_user("frank"))
    assert (await _decide(api, cid, "C-BO", user="dan", confirmed=True, comment="x")).status_code == 403


async def test_measured_review_time_is_reported_with_its_basis(api):
    case = await _run(api)
    await _decide(api, case["case_id"], "E-FO", review_seconds=45)
    ov = (await api.get("/api/overview", headers=api.as_user("frank"))).json()
    assert ov["measured_review"]["decisions"] == 1
    assert ov["measured_review"]["median_seconds"] == 45
    assert "on screen" in ov["measured_review"]["basis"]
